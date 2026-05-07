"""build.py — assemble .pptx from moments + media.

Reads $WORKDIR/moments.json and media from $WORKDIR/media/.
Writes the deck to $WORKDIR/highlights.pptx.
Main process then prompts user where to save it.
"""

import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x21, 0x21, 0x21)
MUTED = RGBColor(0x88, 0x88, 0x88)


def emit(event):
    print(json.dumps(event), flush=True)


def parse_ts(ts):
    parts = [float(p) for p in str(ts).strip().split(":")]
    if len(parts) == 2:
        return parts[0] * 60 + parts[1]
    return parts[0] * 3600 + parts[1] * 60 + parts[2]


def fmt_ts(seconds):
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:d}:{s:02d}"


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK, font="Calibri", italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return box


def build(moments, media_dir, video_title, youtube_url, out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    add_text(s, Inches(0.7), Inches(2.8), Inches(12), Inches(2),
             video_title, size=48, bold=True, color=WHITE, font="Georgia")
    add_text(s, Inches(0.7), Inches(4.7), Inches(12), Inches(0.5),
             f"Notable moments • {len(moments)} highlights",
             size=18, color=ICE, italic=True)

    # Moment slides
    for i, m in enumerate(moments, 1):
        s = prs.slides.add_slide(blank)

        add_text(s, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.9),
                 f"{i:02d}  {m['title']}",
                 size=32, bold=True, color=NAVY, font="Georgia")

        gif_path = media_dir / f"moment_{i:02d}.gif"
        png_path = media_dir / f"moment_{i:02d}.png"

        if gif_path.exists():
            s.shapes.add_picture(str(gif_path), Inches(0.5), Inches(1.6),
                                 width=Inches(6.0))
        if png_path.exists():
            s.shapes.add_picture(str(png_path), Inches(6.8), Inches(1.6),
                                 width=Inches(6.0))

        add_text(s, Inches(0.5), Inches(6.1), Inches(12.3), Inches(0.9),
                 m["description"], size=14)

        start = parse_ts(m["start"]); end = parse_ts(m["end"])
        add_text(s, Inches(0.5), Inches(6.95), Inches(12.3), Inches(0.4),
                 f"{fmt_ts(start)}–{fmt_ts(end)}  •  {youtube_url}",
                 size=10, color=MUTED)

        # Speaker notes — talking points hidden from the audience.
        notes = (m.get("notes") or "").strip()
        if notes:
            s.notes_slide.notes_text_frame.text = notes

        emit({"type": "log", "message": f"  slide {i+1}: {m['title']}"})

    prs.save(out_path)


def main():
    title = sys.argv[1] if len(sys.argv) > 1 else "Video Highlights"
    youtube_url = sys.argv[2] if len(sys.argv) > 2 else ""

    workdir = Path(os.environ["WORKDIR"])
    moments = json.loads((workdir / "moments.json").read_text())
    out_path = workdir / "highlights.pptx"

    emit({"type": "log", "message": f"Building {len(moments)+1} slides…"})
    build(moments, workdir / "media", title, youtube_url, out_path)
    emit({"type": "result", "data": {"path": str(out_path)}})


if __name__ == "__main__":
    main()
